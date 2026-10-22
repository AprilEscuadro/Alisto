"""Admin > Reports > Export: PDF, Excel, or Word.

Every format has the same content:
  1. Key findings   short sentences that explain the numbers
                    (busiest day, time, week, month, barangay, trend...)
  2. Summary        the stat cards from the Reports page
  3. Peak times     alerts by day of week, time of day, week, month, date
  4. Barangays      the Barangay Summary table
  5. Breakdown      alerts by type and by status, residents with most alerts
  6. Alert list     every alert in the period

Needs: pip install reportlab openpyxl python-docx
"""
import io
from collections import Counter, OrderedDict
from datetime import datetime, timedelta

from flask import Response

import alerts_common as ac

DAYS = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']
TIME_BLOCKS = [  # (start hour, label)
    (0, 'Late night (12 AM - 4 AM)'),
    (4, 'Early morning (4 AM - 8 AM)'),
    (8, 'Morning (8 AM - 12 PM)'),
    (12, 'Afternoon (12 PM - 4 PM)'),
    (16, 'Evening (4 PM - 8 PM)'),
    (20, 'Night (8 PM - 12 AM)'),
]
FORMATS = {
    'pdf': ('application/pdf', 'pdf'),
    'excel': ('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'xlsx'),
    'word': ('application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'docx'),
}
ALERT_LIST_LIMIT = 500  # PDF / Word only; Excel always has every alert


# ============================================================
# ANALYSIS
# ============================================================

def _pct(n, total):
    return round(n / (total or 1) * 100, 1)


def _week_label(monday):
    sunday = monday + timedelta(days=6)
    return f"{monday.strftime('%b %d')} - {sunday.strftime('%b %d, %Y')}"


def _ranked(counter, total, order=None):
    """[(label, count, pct)] most first (or in a fixed order)."""
    keys = order if order else [k for k, _ in counter.most_common()]
    return [(k, counter.get(k, 0), _pct(counter.get(k, 0), total)) for k in keys]


def analyze(rows, barangay_rows):
    total = len(rows)
    moments = [(r, ac.local(r['created_at'])) for r in rows]
    moments = [(r, m) for r, m in moments if m]

    by_day = Counter(DAYS[m.weekday()] for _, m in moments)
    by_block = Counter(TIME_BLOCKS[m.hour // 4][1] for _, m in moments)
    by_week = Counter((m.date() - timedelta(days=m.weekday())) for _, m in moments)
    by_month = Counter(m.strftime('%Y-%m') for _, m in moments)
    by_date = Counter(m.date() for _, m in moments)
    by_status = Counter(r['status_label'] for r in rows)
    by_resident = Counter(r['name'] for r in rows if r.get('name'))

    weeks = [(_week_label(w), c, _pct(c, total)) for w, c in
             sorted(by_week.items(), key=lambda kv: (-kv[1], kv[0]))]
    months = [(datetime.strptime(k, '%Y-%m').strftime('%B %Y'), c, _pct(c, total))
              for k, c in sorted(by_month.items(), key=lambda kv: (-kv[1], kv[0]))]
    dates = [(d.strftime('%a, %b %d, %Y'), c, _pct(c, total))
             for d, c in sorted(by_date.items(), key=lambda kv: (-kv[1], kv[0]))]

    a = {
        'total': total,
        'by_day': _ranked(by_day, total, DAYS),
        'by_block': _ranked(by_block, total, [b[1] for b in TIME_BLOCKS]),
        'weeks': weeks,
        'months': months,
        'dates': dates[:10],
        'by_status': _ranked(by_status, total),
        'top_residents': _ranked(by_resident, total)[:10],
    }
    a['findings'] = _findings(a, rows, moments, by_day, by_block, barangay_rows)
    return a


def _top(counter):
    """(label, count) of the biggest, or None. Ties -> all tied labels joined."""
    if not counter or not max(counter.values()):
        return None
    best = max(counter.values())
    names = [k for k, v in counter.items() if v == best]
    return ' and '.join(names), best


def _findings(a, rows, moments, by_day, by_block, barangay_rows):
    total = a['total']
    if not total:
        return ['No emergency alerts were recorded in this period.']

    out = []
    day = _top(by_day)
    if day:
        out.append(f"{day[0]} had the most emergencies: {day[1]} alert(s), "
                   f"{_pct(day[1], total)}% of the period.")
    block = _top(by_block)
    if block:
        out.append(f"Peak time of day is {block[0]} with {block[1]} alert(s) "
                   f"({_pct(block[1], total)}%).")
    if a['weeks']:
        w = a['weeks'][0]
        out.append(f"Busiest week was {w[0]} with {w[1]} alert(s).")
    if len(a['months']) > 1:
        m = a['months'][0]
        out.append(f"Busiest month was {m[0]} with {m[1]} alert(s).")
    elif a['months']:
        out.append(f"All alerts in this period happened in {a['months'][0][0]}.")
    if a['dates']:
        d = a['dates'][0]
        out.append(f"Single busiest day was {d[0]} with {d[1]} alert(s).")

    # Trend: first half of the period vs second half.
    if len(moments) >= 4:
        times = sorted(m for _, m in moments)
        middle = times[0] + (times[-1] - times[0]) / 2
        first = sum(1 for t in times if t < middle)
        second = len(times) - first
        if first and second != first:
            change = round((second - first) / first * 100)
            word = 'increased' if change > 0 else 'decreased'
            out.append(f"Alerts {word} by {abs(change)}% in the second half of the period "
                       f"({first} -> {second}).")

    places = [b for b in barangay_rows if b['alerts']]
    if places:
        b = places[0]
        out.append(f"{b['name']} has the most alerts ({b['alerts']}, "
                   f"{_pct(b['alerts'], total)}%), with a {b['resolved_pct']}% resolution rate.")

    resolved = sum(1 for r in rows if r['status'] in ac.RESOLVED_STATUSES)
    still_open = sum(1 for r in rows if r.get('is_open'))
    out.append(f"{resolved} of {total} alerts were resolved ({_pct(resolved, total)}%)"
               + (f"; {still_open} still need a response." if still_open else "."))

    if a['top_residents'] and a['top_residents'][0][1] >= 3:
        name, count, _ = a['top_residents'][0]
        out.append(f"{name} triggered the most alerts ({count}). "
                   f"A check-up visit by the BHW is recommended.")
    return out


# ============================================================
# SHARED TABLE DATA
# ============================================================

def _summary_rows(stats):
    return [
        ('Total Devices', str(stats['total_devices'])),
        ('Online Devices', f"{stats['online_devices']} ({stats['online_pct']}% of total)"),
        ('Total Alerts', str(stats['total_alerts'])),
        ('Resolved Alerts', f"{stats['resolved']} ({stats['resolved_pct']}% resolution rate)"),
        ('Avg. Response Time', stats['avg_response']),
        ('Active Users', str(stats['active_users'])),
    ]


def _count_table(rows):
    return [[label, str(count), f'{pct}%'] for label, count, pct in rows]


def _sections(a, stats, barangay_rows, type_breakdown):
    """[(title, header, rows)] used by PDF and Word."""
    return [
        ('Summary', ['Metric', 'Value'], [list(r) for r in _summary_rows(stats)]),
        ('Alerts by Day of the Week', ['Day', 'Alerts', 'Share'], _count_table(a['by_day'])),
        ('Alerts by Time of Day', ['Time', 'Alerts', 'Share'], _count_table(a['by_block'])),
        ('Alerts by Week (busiest first)', ['Week', 'Alerts', 'Share'], _count_table(a['weeks'])),
        ('Alerts by Month (busiest first)', ['Month', 'Alerts', 'Share'], _count_table(a['months'])),
        ('Top 10 Busiest Days', ['Date', 'Alerts', 'Share'], _count_table(a['dates'])),
        ('Barangay Summary',
         ['Barangay', 'Residents', 'Devices', 'Online', 'Alerts', 'Resolved', 'Avg. Response'],
         [[b['name'], str(b['residents']), str(b['devices']),
           f"{b['online']} ({b['online_pct']}%)", str(b['alerts']),
           f"{b['resolved']} ({b['resolved_pct']}%)", b['avg_response']] for b in barangay_rows]),
        ('Alerts by Type', ['Type', 'Alerts', 'Share'],
         [[t['label'], str(t['count']), f"{t['pct']}%"] for t in type_breakdown]),
        ('Alerts by Status', ['Status', 'Alerts', 'Share'], _count_table(a['by_status'])),
        ('Residents with the Most Alerts', ['Resident', 'Alerts', 'Share'],
         _count_table(a['top_residents'])),
    ]


def _alert_list(rows):
    header = ['Alert ID', 'Resident', 'Barangay', 'Type', 'Status', 'Date & Time', 'Response (min)']
    body = [[r['code'], r['name'], r['barangay'] or '', r['type_label'], r['status_label'],
             r['full_time'], '' if r['response_minutes'] is None else str(round(r['response_minutes']))]
            for r in rows]
    return header, body


# ============================================================
# PDF
# ============================================================

def _pdf(a, rows, stats, barangay_rows, type_breakdown, range_label, generated):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    TableStyle, ListFlowable, ListItem)

    blue = colors.HexColor('#4a86f7')
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=1.6 * cm, rightMargin=1.6 * cm,
                            topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                            title='ALISTO Report')
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle('h1', parent=ss['Title'], textColor=colors.HexColor('#0f1b3d'))
    h2 = ParagraphStyle('h2', parent=ss['Heading2'], textColor=blue, spaceBefore=12)
    small = ParagraphStyle('small', parent=ss['Normal'], fontSize=8, leading=10)
    body = ss['Normal']

    def table(header, data, widths=None):
        cells = [[Paragraph(f'<b>{h}</b>', small) for h in header]]
        cells += [[Paragraph(str(c), small) for c in row] for row in data] or \
                 [[Paragraph('No data', small)] + [''] * (len(header) - 1)]
        t = Table(cells, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#eaf1ff')),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#d5dbe8')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f7f9fc')]),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        return t

    story = [
        Paragraph('ALISTO Activity Report', h1),
        Paragraph(f'Period: <b>{range_label}</b> &nbsp;&nbsp; Generated: {generated}', body),
        Spacer(1, 10),
        Paragraph('Key Findings', h2),
        ListFlowable([ListItem(Paragraph(f, body)) for f in a['findings']],
                     bulletType='bullet', leftIndent=12),
    ]
    for title, header, data in _sections(a, stats, barangay_rows, type_breakdown):
        story += [Paragraph(title, h2), table(header, data)]

    header, data = _alert_list(rows[:ALERT_LIST_LIMIT])
    story += [Paragraph('Alert List', h2)]
    if len(rows) > ALERT_LIST_LIMIT:
        story += [Paragraph(f'Showing the latest {ALERT_LIST_LIMIT} of {len(rows)} alerts. '
                            'Export to Excel for the full list.', small), Spacer(1, 4)]
    story += [table(header, data)]

    def footer(canvas, d):
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(1.6 * cm, 0.9 * cm, f'ALISTO - {range_label}')
        canvas.drawRightString(A4[0] - 1.6 * cm, 0.9 * cm, f'Page {d.page}')

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ============================================================
# WORD
# ============================================================

def _word(a, rows, stats, barangay_rows, type_breakdown, range_label, generated):
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10)

    doc.add_heading('ALISTO Activity Report', 0)
    p = doc.add_paragraph()
    p.add_run('Period: ').bold = True
    p.add_run(range_label)
    p.add_run('    Generated: ').bold = True
    p.add_run(generated)

    doc.add_heading('Key Findings', 1)
    for f in a['findings']:
        doc.add_paragraph(f, style='List Bullet')

    def table(header, data):
        data = data or [['No data'] + [''] * (len(header) - 1)]
        t = doc.add_table(rows=1, cols=len(header))
        t.style = 'Light Grid Accent 1'
        for cell, h in zip(t.rows[0].cells, header):
            cell.text = h
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for row in data:
            cells = t.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)
        for row in t.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9)

    for title, header, data in _sections(a, stats, barangay_rows, type_breakdown):
        doc.add_heading(title, 2)
        table(header, data)

    doc.add_heading('Alert List', 2)
    if len(rows) > ALERT_LIST_LIMIT:
        note = doc.add_paragraph(f'Showing the latest {ALERT_LIST_LIMIT} of {len(rows)} alerts. '
                                 'Export to Excel for the full list.')
        note.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    table(*_alert_list(rows[:ALERT_LIST_LIMIT]))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ============================================================
# EXCEL
# ============================================================

def _excel(a, rows, stats, barangay_rows, type_breakdown, range_label, generated):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment

    bold = Font(bold=True)
    head_fill = PatternFill('solid', fgColor='EAF1FF')
    title_font = Font(bold=True, size=14, color='0F1B3D')
    section_font = Font(bold=True, size=12, color='4A86F7')

    wb = Workbook()

    def header_row(ws, row, values):
        for col, v in enumerate(values, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font, c.fill = bold, head_fill

    def autosize(ws):
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)

    def count_block(ws, start, title, label, data, chart_anchor=None):
        """Title + table (label, alerts, share); optional bar chart. Returns next free row."""
        ws.cell(row=start, column=1, value=title).font = section_font
        header_row(ws, start + 1, [label, 'Alerts', 'Share (%)'])
        r = start + 2
        for name, count, pct in data:
            ws.cell(row=r, column=1, value=name)
            ws.cell(row=r, column=2, value=count)
            ws.cell(row=r, column=3, value=pct)
            r += 1
        if chart_anchor and data:
            chart = BarChart()
            chart.title, chart.y_axis.title = title, 'Alerts'
            chart.legend = None
            chart.height, chart.width = 7, 14
            chart.add_data(Reference(ws, min_col=2, min_row=start + 1, max_row=r - 1),
                           titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=1, min_row=start + 2, max_row=r - 1))
            ws.add_chart(chart, chart_anchor)
        return r + 1

    # --- Summary ---
    ws = wb.active
    ws.title = 'Summary'
    ws['A1'] = 'ALISTO Activity Report'
    ws['A1'].font = title_font
    ws['A2'], ws['B2'] = 'Period', range_label
    ws['A3'], ws['B3'] = 'Generated', generated
    ws['A5'] = 'Key Findings'
    ws['A5'].font = section_font
    r = 6
    for f in a['findings']:
        ws.cell(row=r, column=1, value=f'- {f}').alignment = Alignment(wrap_text=False)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value='Summary').font = section_font
    header_row(ws, r + 1, ['Metric', 'Value'])
    r += 2
    for metric, value in _summary_rows(stats):
        ws.cell(row=r, column=1, value=metric)
        ws.cell(row=r, column=2, value=value)
        r += 1
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 36

    # --- Peak Times (tables + charts) ---
    ws = wb.create_sheet('Peak Times')
    r = count_block(ws, 1, 'Alerts by Day of the Week', 'Day', a['by_day'], 'E1')
    r = count_block(ws, max(r, 17), 'Alerts by Time of Day', 'Time', a['by_block'], f'E{max(r, 17)}')
    r = count_block(ws, max(r, 33), 'Alerts by Month', 'Month', a['months'], f'E{max(r, 33)}')
    r = count_block(ws, max(r, 49), 'Alerts by Week', 'Week', a['weeks'])
    count_block(ws, r, 'Top 10 Busiest Days', 'Date', a['dates'])
    ws.column_dimensions['A'].width = 30

    # --- Barangays ---
    ws = wb.create_sheet('Barangays')
    header_row(ws, 1, ['Barangay', 'Residents', 'Devices', 'Online', 'Online (%)',
                       'Alerts', 'Resolved', 'Resolved (%)', 'Avg. Response'])
    for i, b in enumerate(barangay_rows, 2):
        for col, v in enumerate([b['name'], b['residents'], b['devices'], b['online'],
                                 b['online_pct'], b['alerts'], b['resolved'],
                                 b['resolved_pct'], b['avg_response']], 1):
            ws.cell(row=i, column=col, value=v)
    autosize(ws)

    # --- Breakdown ---
    ws = wb.create_sheet('Breakdown')
    r = count_block(ws, 1, 'Alerts by Type', 'Type',
                    [(t['label'], t['count'], t['pct']) for t in type_breakdown])
    r = count_block(ws, r, 'Alerts by Status', 'Status', a['by_status'])
    count_block(ws, r, 'Residents with the Most Alerts', 'Resident', a['top_residents'])
    ws.column_dimensions['A'].width = 30

    # --- Alerts (every row) ---
    ws = wb.create_sheet('Alerts')
    header, _ = _alert_list([])
    header_row(ws, 1, header)
    for i, x in enumerate(rows, 2):
        for col, v in enumerate([x['code'], x['name'], x['barangay'] or '', x['type_label'],
                                 x['status_label'], x['full_time'],
                                 None if x['response_minutes'] is None
                                 else round(x['response_minutes'])], 1):
            ws.cell(row=i, column=col, value=v)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    autosize(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ============================================================
# ENTRY POINT
# ============================================================

def build(fmt, rows, stats, barangay_rows, type_breakdown, range_label):
    """Flask Response with the report file, or None for an unknown format."""
    if fmt not in FORMATS:
        return None
    now = datetime.now(ac.PH_TIME)
    generated = now.strftime('%b %d, %Y %I:%M %p')
    a = analyze(rows, barangay_rows)
    maker = {'pdf': _pdf, 'excel': _excel, 'word': _word}[fmt]
    data = maker(a, rows, stats, barangay_rows, type_breakdown, range_label, generated)

    mimetype, ext = FORMATS[fmt]
    slug = range_label.lower().replace(' ', '-')
    filename = f"alisto-report-{slug}-{now.strftime('%Y-%m-%d')}.{ext}"
    return Response(data, mimetype=mimetype,
                    headers={'Content-Disposition': f'attachment; filename={filename}'})